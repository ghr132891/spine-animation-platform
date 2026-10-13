"""Build the customer handover from implemented scope and saved evidence."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'delivery';OUT.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0];sec.page_width=Cm(21.59);sec.page_height=Cm(27.94)
sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.9);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
sec.footer_distance=Cm(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
    s=doc.styles[name];s.font.name='Arial Unicode MS';s.font.color.rgb=RGBColor(0,0,0)
    s.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Arial Unicode MS')
    s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.18
    for border in list(s.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    fonts=s.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ('asciiTheme','eastAsiaTheme','hAnsiTheme','cstheme'):fonts.attrib.pop(qn('w:'+key),None)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Title'].font.size=Pt(23);doc.styles['Title'].paragraph_format.space_after=Pt(10)
doc.styles['Subtitle'].font.size=Pt(10.5)
for name,size in [('Heading 1',16),('Heading 2',12)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(10);doc.styles[name].paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Spine 工作台交付说明  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
doc.core_properties.title='Spine 工作台功能与改进交付说明'
doc.core_properties.subject='已实现功能 验证范围 现有问题 后续改进'
doc.core_properties.author='';doc.core_properties.last_modified_by=''

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text,level=1):return doc.add_heading(text,level)
def bullet(text):return p(text,'List Bullet')
def page(title):h(title)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for i,values in enumerate([headers,*rows]):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,text in enumerate(values):
            cell=cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=str(text)
            props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for edge,val in [('top',80),('bottom',80),('left',100),('right',100)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF3');props.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.15
                if j==0 and len(widths)>2:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size=Pt(9.5);r.font.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
            no_split=OxmlElement('w:cantSplit');cell._tc.getparent().get_or_add_trPr().append(no_split)
    p('').paragraph_format.space_after=Pt(0)
    return t

p('Spine 工作台功能与改进交付说明','Title')
p('交付状态记录  2026年10月9日','Subtitle')
p('本次成果包括已有 Spine 的贴图换皮、骨骼与贴图编辑、单图角色行走制作，以及供 AI 调用的自动换皮接口和 Skill。现有 Spine 换皮与 Agent 自动调用流程已完成可用版本；单图行走已能生成资源，但用户仍反馈视觉问题，后续需要继续调整后再验收效果。')
h('功能交付状态')
table(['模块','当前状态','交付内容与范围'],[
 ('贴图换皮','已实现','保留原骨骼和动作，修改部件或图集，预览并导出。'),
 ('骨骼与贴图','已实现','人工编辑骨架、附件和 IK，替换贴图，导出新的 JSON 资源。'),
 ('单图行走','已实现流程\n视觉待改善','拆件、关节定位、组装校准、生成 idle 和 walk；效果仍需调整。'),
 ('AI 自动换皮','已实现','目标风格生成提示词，亿创生图，自动替换、校验和导出。'),
 ('批量调用','已支持部件批量\n多角色按任务调用','一套资源批量处理部件；多套资源由 Agent 分别建任务。'),
 ('局域网与 Skill','已实现','监听 0.0.0.0:8780；提供 OpenAPI、CLI 和 spine-ai-reskin Skill。'),
],[3.1,3.6,10.3])
h('当前访问入口')
p('本机：http://127.0.0.1:8780/\n局域网：http://172.16.7.76:8780/\n另一网卡：http://172.16.64.213:8780/')
p('以上局域网地址为本次部署记录，IP 变化后使用服务器实际地址。页面和接口已通过本机网卡地址验证；其他设备的网络连通性需在对应网络环境复验。')
h('交付口径')
p('可交付工作台、自动调用能力及可运行资源样例。单图行走按试用能力交付，并附后续问题清单。当前不承诺任意角色的一键高质量重绘，也未实现任意骨架之间的自动动作迁移。')

page('已有 Spine 的换皮与编辑')
h('保持骨骼的贴图换皮',2)
p('导入 JSON 或 SKEL、atlas、全部 PNG，或一套 ZIP 资源。支持子目录、多页图集、裁切、正交旋转及显式 PMA。PMA 表示预乘透明度，需遵循原资源声明。')
bullet('提供部件单选和多选、原部件下载、图集页下载；可调整色相、饱和度和明度，上传同尺寸部件图或替换整页图集。')
bullet('预览支持 Spine 3.8 和 4.2，切换动作、皮肤、速度、骨骼及整体镜像；朝向翻转会同时翻转脚掌。')
bullet('导出保留原骨架和 atlas 的字节，锁定 PNG 尺寸与 alpha，附文件校验报告。沿用已有 mesh、权重、deform、IK、events 与动作。')
h('骨骼和贴图编辑',2)
p('可从已有项目创建独立编辑方案，调整骨骼位置、旋转、缩放、长度、名称和父子关系；新增骨骼，删除符合引用约束的叶骨骼。名称变更会同步相关引用。')
bullet('支持任意尺寸及透明度的部件 PNG、更换插槽绑定、增删普通图片附件、编辑附件位置和旋转，并重打多页图集。')
bullet('支持编辑双骨 IK 的链条、目标、混合、柔化和弯曲方向；支持导入使用当前图集、相同版本的完整骨架 JSON。')
bullet('选中骨骼后显示名称、关节点、骨长、影响贴图及子骨骼关系；可在画布选中、定位和拖动。编辑有未保存预览、应用及取消操作。')
p('该入口支持 3.8 和 4.2 JSON；SKEL 需先转换为匹配版本的 JSON。编辑后沿用局部动作曲线，比例或层级变化仍需逐动作检查；完全换成不同骨架时不会自动重定向动作。')
h('历史记录与任务持久化',2)
p('项目、原图、编辑版本和 AI 任务持久保存在服务器。支持命名版本、撤销、恢复及重新打开；恢复会另存版本，保留已有后续记录。页面 URL 同步项目编号，便于刷新和重开。')
p('亿创任务提交前保存编号，关闭页面后继续查询；历史查询不会重新提交生图。查询失败有退避和上限，已有结果可再次查询和下载。旧系统中更早已被覆盖且未保存的状态不能追溯恢复。')
h('导出格式',2)
p('第一步保留原 JSON 或 SKEL；第二步输出 JSON、atlas、图集、部件 PNG 和校验报告。第二步不生成原生 .spine 工程或新的二进制 SKEL。')

page('单张角色图制作行走')
p('支持上传 PNG、JPG 或 WebP，或通过亿创生成全身角色图。推荐关节点按轮廓比例估计，可拖动或填写坐标；支持去背景、腿间背景清理、部件多边形修正及原图位置高亮。')
p('可建立人形双足骨架，生成 idle 与原地循环 walk，并调整周期、步幅、抬脚高度和朝向。输出 Spine 4.2 JSON、atlas、图集、部件、原图及关节配置，可继续进入骨骼与贴图编辑。')
h('已实施的修正',2)
bullet('针对撕裂，增加共享网格与连续权重；针对粘连，按身体、左右手臂和左右腿隔离蒙皮。')
bullet('收紧身体、背包和脚的拆分边界，清理邻近部件像素，并提供部件定位与高亮。')
bullet('接入身体与背包、手臂、双腿的独立部件参考生成；默认只补遮挡区域，保留原图可见像素和关节位置。')
bullet('为当前机器人调整鞋口与小腿绑定；用原图近侧鞋的完整鞋跟补齐后侧缺失处。保留各版方案供比较和恢复。')
pic=doc.add_picture(str(ROOT/'evidence/single-image/comparison.png'),width=Cm(14.5))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER;doc.paragraphs[-1].paragraph_format.keep_with_next=True
p('图示为原角色和行走采样的离线渲染，用于展示实现流程；不代表用户已接受当前视觉效果。').paragraph_format.space_after=Pt(6)
h('仍需改善的效果',2)
p('用户此前反馈仍有问题。本次增加原图对照、部件校准和拆分检查，便于继续调整；尚未完成视觉效果验收。后续重点复验身体与背包分层、被遮挡的手臂和脚、补全后的原角色一致性，以及迈步时的鞋口连接、落地和穿插。已有像素和循环校验通过，不等于这些视觉问题已全部解决。')
p('当前适用于站立且四肢可辨的人形双足角色。任意姿态识别、完整遮挡区域恢复及复杂角色通用拆件未达到自动化交付标准。')

page('参考项目流程的落地改进')
p('参考 GenielabsOpenSource/spine-animation-ai 的拆件摆放与可逆校准流程，新增“组装校准”标签页。先对照原图检查部件，再制作行走，减少位置和前后层级偏差被带入动画。实现独立编写，未复制参考项目的代码或素材。')
pic=doc.add_picture(str(ROOT/'evidence/assembly-review/browser-assembly.jpg'),width=Cm(14.5))
doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER;doc.paragraphs[-1].paragraph_format.keep_with_next=True
p('实际浏览器界面：原图与组装并排，选中身体部件有定位轮廓；下方提供校准表单。').paragraph_format.space_after=Pt(6)
h('本次新增能力',2)
bullet('可选中并拖动部件，填写横移、纵移、旋转和缩放；提供原图叠加、单部件查看、前后移层、单部件复位及全部复位。')
bullet('连通像素块检查提示断开连接或邻近部件像素；补全超过 20% 提示对照原图。检查不会删除零碎装饰，也不能保证部件语义正确。')
bullet('校准随行走历史保存和恢复，导出新增 assembly-review.json；零调整不重采样，调整仅改变对应贴图几何和插槽顺序，骨骼与动作曲线保持。')
h('供 AI 读取和修正',2)
p('GET /api/characters/{id}/review 返回原图尺寸、部件范围、锚点、调整和检查结果；POST 同一路径保存 revision 和 part_review，包含 adjustments 与 draw_order。GET /assembly 返回组装透明 PNG。使用原图像素坐标，Y 向下、顺时针为正。')
p('本次未提交新的 AI 生图。参考项目的 SIFT/RANSAC 自动配准、自动遮挡层级推断及更多动作预设尚未接入；缺失像素、角色一致性及行走效果仍需改善。')

page('AI 调用和批量能力')
p('输入原 Spine 资源或项目编号及目标风格，服务自动生成部件提示词、调用亿创、收集图片、按原轮廓贴回、校验并导出独立新项目。源项目不会被任务覆盖。')
h('生成路线',2)
table(['路线','当前能力与条件'],[
 ('material 默认','生成表面材质，结合原图明暗贴回；保留原暗色线条、轮廓和几何，适合颜色与材质风格变化。'),
 ('reference','按部件参考画布重绘；需为每个部件提供对应的 512×512 画布公共 HTTPS 地址。当前不自动上传源素材。'),
],[3,14])
p('提示词根据部件名称、尺寸、原色统计和目标风格构造，可在计划中查看实际模型输入。当前不包含视觉语义模型识别或额外的 LLM 提示词优化调用。')
h('接口与客户端',2)
table(['操作','接口'],[
 ('查看调用规范','GET /api/automation/openapi.json'),
 ('生成计划 不提交生图','POST /api/automation/reskin/plan'),
 ('提交任务 查询任务列表','POST 或 GET /api/automation/reskin/tasks'),
 ('查询 下载 恢复 取消','在 tasks/{id} 下查询，或调用 download、resume、cancel'),
],[5.1,11.9])
p('POST 使用 JSON 和 X-Spine-Local: 1。亿创凭据在服务器端配置并验证，以 0600 权限保存，不返回调用方，也不写入生成资源包。')
p('CLI 为 tools/spine_agent.py；Skill 为 spine-ai-reskin，已安装到本机个人 Skills 目录，并提供包含客户端和 OpenAPI 的可分享 ZIP。局域网调用指定 --base http://服务器IP:8780。')
h('批量与恢复边界',2)
bullet('同一套 Spine 支持全部或指定部件批量处理。默认上限 32 次，可配置为 1–128；一张部件对应一次生图。')
bullet('多套角色可由 Agent 分别建立独立任务。当前没有一次上传多角色的统一批次接口、目录自动遍历命令或批次总看板。')
bullet('同一 nonce 和相同参数复用任务，避免重复调用。提交结果不明时先核实已有 gen_id，不能自动盲目重新生成。')
bullet('查询有上限，支持恢复、下载和取消后续处理。客户端超时不取消服务器任务；取消也不代表已取消亿创端的已有生成。')
p('已完成真实亿创单部件自动换皮和导出验证；整套部件批量通过模拟测试。尚未完成多角色真实生图的规模、成本和视觉一致性验收。')

page('后续改善问题与建议顺序')
p('下面是建议的改进顺序，不代表已排期或已承诺交付日期。优先完成用户已指出的单图视觉问题，再扩大自动生成与批量处理范围。')
table(['优先级','问题与当前状态','建议改进与验收目标'],[
 ('P1','单图分层与遮挡补全仍有视觉问题','继续修正身体、背包、手臂和脚的分层；检查摆臂与迈步全周期，无缺块、明显粘连或露缝。'),
 ('P1','补全部件可能偏离原角色','优先保留可见像素，改善隐藏区域的色彩、材质与比例衔接；由用户确认角色一致性。'),
 ('P1','行走质量仍待确认','检查鞋口连接、脚底接触、步幅、穿插与重心；统一地面基准，完成实际预览效果验收。'),
 ('P1','当前视觉检查覆盖不完整','补齐真实浏览器操作与逐动作检查，增加多种角色及姿态样例；区分文件通过与效果通过。'),
 ('P2','参考图重绘依赖外部 HTTPS 画布','接入稳定的素材上传和 URL 管理；完成整套部件参考重绘的一致性验收。'),
 ('P2','提示词依赖名称与规则模板','增加可选视觉描述和风格约束，对无语义命名、左右配对及整套材质一致性做检查。'),
 ('P2','多角色批量缺少统一管理','增加目录遍历、批次接口、进度汇总、生成预算、受控并发及失败项重做；避免重做成功项。'),
 ('P2','局域网多人使用能力有限','增加身份与项目权限、操作审计、备份恢复及服务运维；固定操作头不能等同于用户认证。'),
 ('P3','任意骨架动作迁移尚未接入','研究关节匹配、曲线转换、比例补偿和重新蒙皮；使用不同骨架验证脚掌接地和动作保持。'),
 ('P3','复杂姿态及更多格式仍有限','扩展姿态识别、非标准角色和版本兼容；明确 .spine 工程及二进制转换所需工具链。'),
],[1.5,5.3,10.2])
h('交付时应保留的说明',2)
p('现有骨架编辑属于人工调整后继续沿用局部动画；不能描述为任意骨架自动重定向。默认 AI 自动换皮是材质路线；不能描述为已完成整套角色高保真自动重画。单图行走可展示流程和资源，但当前视觉效果仍需后续调整。')

page('验证记录与交付文件')
h('已有验证结果',2)
table(['验证项','记录结果','验证边界'],[
 ('单元测试','本次全量 56 项通过','覆盖贴图、历史、骨骼、单图与自动任务；不是视觉通过率。'),
 ('自动换皮实测','1 次真实亿创部件任务完成并导出','示例替换身体；不是整套角色或多角色规模实测。'),
 ('动作保持','idle 和 walk 各采样 49 帧\n顶点最大差异为 0','对该换皮示例，骨架及 atlas 字节保持一致。'),
 ('骨骼编辑兼容','3.8 和 4.2 示例及 mesh 和 IK 加载','覆盖实际网格样例；新比例与层级仍需视觉检查。'),
 ('单图原貌保持','可见像素保留 原关节恢复\n循环误差为 0','不能据此排除补绘细节和步态上的视觉问题。'),
 ('组装校准','idle 和 walk 各采样 49 帧\n其他部件误差与循环误差均为 0','骨骼曲线保持；浏览器已检查页面加载、选中和排版。'),
 ('局域网与 Skill','网卡 IP 页面与 API 可用\nSkill 客户端及下载验证通过','跨设备实机网络复验、多用户权限仍需补充。'),
],[3.2,6.4,7.4])
p('验证以自动化测试、DOM 与 HTTP 调用、官方 Spine runtime 采样及离线渲染为主。组装校准页已做实际浏览器加载、部件选择和排版检查；全部编辑动作及其他模块的完整浏览器验收仍需补齐。')
h('成果位置',2)
p('项目根目录：/Users/tu/Auto Creating/spine-reskin').paragraph_format.keep_with_next=True
table(['内容','相对项目根目录的位置'],[
 ('平台源码与启动','app/、web/、tools/、tests/、Start.command'),
 ('使用及接口文档','README.md、agent/README.md、agent/openapi.json'),
 ('Skill 源码及分享包','skills/spine-ai-reskin/\nagent/spine-ai-reskin-skill.zip'),
 ('AI 自动换皮样例和报告','evidence/automation/live-spine.zip\nevidence/automation/acceptance.json'),
 ('单图行走样例和报告','evidence/single-image/preserved-fill/character-preserved.zip\nevidence/single-image/preserved-fill/acceptance.json'),
 ('组装校准改进与报告','evidence/assembly-review/acceptance.json\nevidence/assembly-review/runtime-validation.json'),
 ('其他复验记录','evidence/rerig/、evidence/lan-validation.json\nevidence/automation/skill-validation.json'),
],[5.1,11.9])
h('接收方复验建议',2)
p('先启动平台并导入标准 Spine 样例，完成修改、预览、保存历史和导出。随后用 Agent 的 plan 查看提示词和数量，再按授权范围提交一个小规模任务，确认恢复和下载。单图角色的分层、原貌与行走效果单独组织视觉验收。')
p('运行环境需要 Python 3.10 以上、Pillow、numpy；CLI 客户端仅使用 Python 标准库。项目数据在 storage/，移交测试资源时另行管理服务器凭据，分享包不包含工作台 Key。保留随运行库提供的许可文件。')

target=OUT/'Spine工作台功能与改进交付说明.docx';doc.save(target)
assert json.loads((ROOT/'evidence/assembly-review/acceptance.json').read_text())['tests']['passed']==56
assert json.loads((ROOT/'evidence/single-image/preserved-fill/acceptance.json').read_text())['all_source_pixels_preserved']
print(target)
