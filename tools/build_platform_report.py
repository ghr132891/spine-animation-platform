"""Create an illustrated platform report from saved implementation evidence.

Uses recorded results only. Does not generate images, submit AI tasks, change
the workbench, or replace the previous handover document.
"""
from pathlib import Path
import hashlib
import json
import shutil

import numpy as np
from PIL import Image, ImageDraw, ImageFont, ImageOps
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BATCH = ROOT / 'delivery/flambe-ice-blue-white-20261009'
TRIAL = ROOT.parent / 'spine-upstream-trial'
OUT = ROOT / 'delivery/platform-report-20261009'
ASSETS = OUT / 'assets'
ASSETS.mkdir(parents=True, exist_ok=True)
FONT = '/System/Library/Fonts/STHeiti Medium.ttc'
INK = '#18324b'
BLUE = '#408ebd'
PALE = '#edf4fa'
DOC_OUT = OUT / 'SpineAI工作台平台功能与不足图文总结_反馈更新版.docx'


def font(n): return ImageFont.truetype(FONT, n)


def fit_picture(path, size, bg='#e5ebf3'):
    im = Image.open(path).convert('RGBA')
    im.thumbnail(size, Image.Resampling.LANCZOS)
    layer = Image.new('RGB', size, bg)
    layer.paste(im, ((size[0]-im.width)//2, (size[1]-im.height)//2), im)
    return layer


def diagrams():
    for source, name in [
        ('/var/folders/7b/dbck7xzd1r3ct6h0m3h9w9f40000gn/T/codex-clipboard-1f313599-34ac-4239-b161-11e1725909d7.png', 'user-accepted-walk.png'),
        ('/var/folders/7b/dbck7xzd1r3ct6h0m3h9w9f40000gn/T/codex-clipboard-7edef30a-7ad1-4021-bd7b-849844e6bc5d.png', 'user-matrix-prompt.png'),
    ]:
        if not (ASSETS / name).exists(): shutil.copy2(source, ASSETS / name)
    # Diagram 1 distinguishes the three production routes.
    im = Image.new('RGB', (1600, 680), 'white'); d = ImageDraw.Draw(im)
    rows = [
        ('已有 Spine 换皮', ['原 JSON / atlas / PNG', '选择部件与目标风格', '修改贴图并保留动作', '校验与导出 ZIP']),
        ('骨骼与贴图编辑', ['已有 Spine JSON', '创建独立编辑方案', '调整骨骼 附件 IK', '重打图集与动作检查']),
        ('单图行走制作', ['角色原图或 AI 全身图', '拆件 补全 组装校准', '人形绑骨与行走预设', '预览与视觉验收']),
    ]
    for row, (label, nodes) in enumerate(rows):
        y = 40 + row * 215
        d.text((20, y), label, font=font(32), fill=INK)
        for i, node in enumerate(nodes):
            x = 25 + i * 397
            d.rounded_rectangle((x, y+60, x+353, y+145), 12, fill=PALE, outline='#b8cddd', width=2)
            d.text((x+18, y+88), node, font=font(25), fill=INK)
            if i < 3:
                d.line((x+359, y+103, x+387, y+103), fill=BLUE, width=4)
                d.polygon([(x+387, y+103), (x+377, y+96), (x+377, y+110)], fill=BLUE)
    im.save(ASSETS / 'routes.png')

    # Diagram 2 documents request identity and result collection, not quality.
    im = Image.new('RGB', (1600, 440), 'white'); d = ImageDraw.Draw(im)
    nodes = ['输入资源与风格', 'plan 核对部件与预算', 'run 保存任务编号', '收集候选并校验', '独立结果项目与 ZIP']
    for i, node in enumerate(nodes):
        x = 20 + i * 316
        d.ellipse((x+120, 18, x+160, 58), fill=BLUE)
        d.text((x+132, 25), str(i+1), font=font(24), fill='white')
        d.rounded_rectangle((x, 87, x+290, 177), 9, fill=PALE, outline='#b8cddd', width=2)
        d.text((x+14, 119), node, font=font(23), fill=INK)
        if i < 4:
            d.line((x+294, 132, x+310, 132), fill=BLUE, width=3)
            d.polygon([(x+310,132),(x+303,126),(x+303,138)], fill=BLUE)
    d.text((20, 225), '查询超时 → 保留同一任务继续查询       提交不明 → 核对已有生成编号', font=font(29), fill=INK)
    d.text((20, 290), '生成失败 → 保留候选与原因       重做需要授权       文件通过后仍需视觉检查', font=font(29), fill=INK)
    d.text((20, 365), '本批另加本地配色清理与拟色补齐  原自动任务状态不改写', font=font(25), fill='#61758a')
    im.save(ASSETS / 'automation.png')

    report = json.loads((BATCH / 'final-report.json').read_text())
    counts = [('成功候选', report['completed_candidates'], '#408ebd'), ('失败候选', report['failed_candidates'], '#c47759'), ('未提交部件', report['unsubmitted'], '#b1bcc8')]
    im = Image.new('RGB', (1600, 410), 'white'); d = ImageDraw.Draw(im)
    d.text((10, 15), '冰雪蓝白批次  89 个计划部件', font=font(34), fill=INK)
    x = 15
    for label, value, color in counts:
        width = 1560 * value / 89
        d.rectangle((x, 95, x+width, 200), fill=color)
        d.text((x+14, 126), str(value), font=font(38), fill='white' if value != 16 else INK)
        x += width
    for i, (label, value, color) in enumerate(counts):
        x = 20+i*520
        d.rectangle((x, 246, x+24, 270), fill=color)
        d.text((x+40, 240), f'{label}  {value}', font=font(30), fill=INK)
    d.text((20, 324), '实际提交 73 次   成功候选 60 / 73   AI 整套完成 3 / 8   最终拟色交付 8 / 8', font=font(29), fill=INK)
    im.save(ASSETS / 'batch-counts.png')

    # Original, flawed AI result, final fit. Keep the actual user screenshot.
    screenshot = Path('/var/folders/7b/dbck7xzd1r3ct6h0m3h9w9f40000gn/T/codex-clipboard-c9d67ac6-ffb5-4981-8dcc-5460ac717beb.png')
    if not (ASSETS / 'user-flawed-workbench.png').exists(): shutil.copy2(screenshot, ASSETS / 'user-flawed-workbench.png')
    comparison = Image.open(BATCH / 'evidence/zhujue-comparison.jpg')
    pieces = [comparison.crop((4,40,314,440)), Image.open(ASSETS / 'user-flawed-workbench.png').crop((996,315,1367,845)), comparison.crop((326,40,636,440))]
    im = Image.new('RGB', (1350, 625), '#f4f8fc'); d = ImageDraw.Draw(im)
    for i, (piece, text) in enumerate(zip(pieces, ['原角色', '清理前 AI 结果', '最终配色拟色版'])):
        piece.thumbnail((440, 550), Image.Resampling.LANCZOS)
        layer = Image.new('RGB', (440,550), '#e5ebf3')
        layer.paste(piece, ((440-piece.width)//2,(550-piece.height)//2))
        im.paste(layer, (i*450+5,60)); d.text((i*450+18,17), text, font=font(29), fill=INK)
    im.save(ASSETS / 'identity-comparison.png')

    # Show the observed rectangle overlap and actual mesh UV ownership.
    source = Image.open(BATCH / 'inputs/xiaochushen/xiaochushen.png').convert('RGBA')
    base = Image.new('RGB', source.size, '#e5ebf3'); base.paste(source, mask=source.getchannel('A'))
    rect = base.copy(); draw = ImageDraw.Draw(rect)
    draw.rectangle((406,87,605,236), outline='#dc7658', width=4)
    draw.rectangle((578,117,734,366), outline='#338dbd', width=4)
    overlay = Image.new('RGBA', source.size); draw = ImageDraw.Draw(overlay)
    uv = json.loads((BATCH / 'evidence/xiaochushen-uv.json').read_text())
    for item in uv:
        if item['name'] not in ('xiabanshen', 'zuodabi'): continue
        points = np.array(item['uv']).reshape(-1,2) * np.array(source.size)
        color = (220,118,88,100) if item['name']=='xiabanshen' else (51,141,189,100)
        for tri in np.array(item['triangles']).reshape(-1,3): draw.polygon([tuple(p) for p in points[tri]], fill=color)
    masked = Image.alpha_composite(base.convert('RGBA'), overlay).convert('RGB')
    im = Image.new('RGB', (1350, 625), '#f4f8fc'); d = ImageDraw.Draw(im)
    for i, (pic, title) in enumerate([(rect,'atlas 矩形范围交叠'),(masked,'原 mesh UV 彩色范围示意')]):
        pic = pic.crop((385,72,755,390)); pic = ImageOps.contain(pic,(650,540),Image.Resampling.LANCZOS)
        im.paste(pic,(i*675+(675-pic.width)//2,65));d.text((i*675+18,18),title,font=font(28),fill=INK)
    im.save(ASSETS / 'uv-overlap.png')

    # Use four sampled GIF frames as a static motion strip.
    gif = Image.open(ROOT / 'evidence/single-image/preserved-fill/walk.gif')
    strip = Image.new('RGB', (1600, 610), '#f4f8fc'); d = ImageDraw.Draw(strip)
    for i, fraction in enumerate([0,.25,.5,.75]):
        gif.seek(min(gif.n_frames-1, int(gif.n_frames*fraction)))
        pic = ImageOps.contain(gif.convert('RGB'),(390,540),Image.Resampling.LANCZOS)
        strip.paste(pic,(i*400+(400-pic.width)//2,60));d.text((i*400+16,15),f'周期 {int(fraction*100)}%',font=font(27),fill=INK)
    strip.save(ASSETS / 'robot-walk.png')

    # Original project's AI sheet alongside the Agent-assisted 16-part sheet.
    left = fit_picture(TRIAL / 'site/media/generated-sheet.png', (650,650), '#ffffff')
    right = fit_picture(TRIAL / 'site/media/assisted-parts-sheet.png', (650,650), '#ffffff')
    im = Image.new('RGB',(1350,715),'#f4f8fc');d=ImageDraw.Draw(im)
    im.paste(left,(5,60));im.paste(right,(695,60))
    d.text((18,17),'生成图 连通域仅拆出 5 块',font=font(26),fill=INK)
    d.text((710,17),'Agent 辅助切为 16 个部件',font=font(26),fill=INK)
    im.save(ASSETS / 'upstream-parts.png')

    # One side-by-side full figure is more readable than the original page UI.
    bone = fit_picture(ROOT / 'evidence/rerig/robot42/bone-inspection.png', (560,670))
    other = fit_picture(ROOT / 'evidence/rerig/mesh38/bone-inspection.png', (560,670))
    im = Image.new('RGB',(1160,735),'#f4f8fc');d=ImageDraw.Draw(im)
    im.paste(bone,(10,55));im.paste(other,(590,55))
    d.text((18,14),'4.2 选中骨骼与影响部件',font=font(25),fill=INK)
    d.text((599,14),'3.8 mesh 影响范围',font=font(25),fill=INK)
    im.save(ASSETS / 'bone-inspection.png')


def init_doc():
    doc = Document(); section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = Inches(.70); section.bottom_margin = Inches(.70)
    section.left_margin = Inches(.76); section.right_margin = Inches(.76)
    section.footer_distance = Inches(.29)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet','Caption']:
        style = doc.styles[name]
        style.font.name = 'Arial Unicode MS'; style.font.color.rgb = RGBColor(0,0,0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('ascii','hAnsi','eastAsia'): fonts.set(qn('w:'+key),'Arial Unicode MS')
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'): fonts.attrib.pop(qn('w:'+key),None)
        for border in list(style.element.xpath('.//w:pBdr')): border.getparent().remove(border)
        style.paragraph_format.space_after = Pt(6); style.paragraph_format.line_spacing = 1.12
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Title'].font.size = Pt(22); doc.styles['Title'].paragraph_format.space_after = Pt(11)
    doc.styles['Subtitle'].font.size = Pt(10.5)
    doc.styles['Subtitle'].font.italic = False
    for name, size in [('Heading 1',16),('Heading 2',12)]:
        s = doc.styles[name]; s.font.size = Pt(size)
        s.paragraph_format.space_before = Pt(9); s.paragraph_format.space_after = Pt(7)
    doc.styles['Caption'].font.size = Pt(9)
    doc.styles['Caption'].font.color.rgb = RGBColor.from_string('566779')
    doc.styles['Caption'].paragraph_format.space_after = Pt(9)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('Spine AI 工作台功能与不足总结   '); r.font.size = Pt(8);r.font.color.rgb=RGBColor.from_string('666666')
    field = OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.core_properties.title = 'Spine AI 工作台平台功能与不足总结'
    doc.core_properties.subject = '平台功能 制作与优化过程 真实批量案例 当前不足与改进建议'
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    return doc


def build():
    diagrams()
    doc = init_doc(); figure_number = 0
    def p(text, style=None): return doc.add_paragraph(text,style)
    def h(text,level=1): return doc.add_heading(text,level)
    def bullet(text): return p(text,'List Bullet')
    def page(title):
        heading = h(title)
        heading.paragraph_format.page_break_before = True
    def figure(path,caption,width=6.75):
        nonlocal figure_number
        figure_number += 1
        para = p(''); para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(3); para.paragraph_format.keep_with_next = True
        run=para.add_run(); inline=run.add_picture(str(path),width=Inches(width))
        inline._inline.docPr.set('descr',caption)
        para=p(f'图 {figure_number}  {caption}','Caption');para.alignment=WD_ALIGN_PARAGRAPH.LEFT
    def table(headers,rows,widths,small=10,compact=False):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for i,w in enumerate(widths):t.columns[i].width=Inches(w)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i, values in enumerate([headers,*rows]):
            cells=t.rows[0].cells if i==0 else t.add_row().cells
            nosplit=OxmlElement('w:cantSplit');cells[0]._tc.getparent().get_or_add_trPr().append(nosplit)
            for j,text in enumerate(values):
                cell=cells[j];cell.width=Inches(widths[j]);cell.text=str(text);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                for edge,val in [('top',60 if compact else 95),('bottom',60 if compact else 95),('left',110),('right',110)]:
                    e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
                props.append(margins)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'214B6B' if i==0 else ('F1F6FA' if i%2==0 else 'FFFFFF'));props.append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1.02 if compact else 1.12
                    if j==0 and widths[j]<1.3:para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for r in para.runs:r.font.size=Pt(small);r.font.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
        p('').paragraph_format.space_after=Pt(0)
        return t
    def hyperlink(label,url):
        para=p(''); link=OxmlElement('w:hyperlink')
        from docx.opc.constants import RELATIONSHIP_TYPE
        link.set(qn('r:id'),para.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
        run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245E85');props.append(color);run.append(props)
        text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);para._p.append(link)

    # 1 Overview, with the main outcome and the current delivery boundary.
    p('Spine AI 工作台\n平台功能与不足总结','Title')
    p('制作与优化过程及真实批量案例   2026年10月9日','Subtitle')
    p('平台已具备已有 Spine 换皮、骨骼与贴图编辑、单图人形行走制作，以及供 AI 调用的接口和 Skill。现阶段适合已有资源的颜色与材质调整、人工辅助骨架编辑和基础动画制作。高保真部件重绘、复杂遮挡补全及任意骨架的动作迁移仍需继续完善。')
    p('本次 Flambé 冰雪蓝白案例最终交付 8 套完整资源，并保留原肤色、五官和动作。AI 自动任务仅 3 套完整成功；最终 8 套均经过原图明暗拟色清理，其中 5 套使用成功候选的共同配色补齐。工作台已同步最终资源，预览和导出贴图与交付包一致。[S3][S4]')
    figure(BATCH/'evidence/workbench-corrected.png','最终修正版已同步回工作台 原角色与蓝白结果并排显示 [S4]',6.7)
    h('当前成果应如何理解',2)
    bullet('已实现：导入、编辑、预览、历史、导出、AI 任务持久化与 Agent 调用。')
    bullet('已验证案例：8 套资源、17 个原版动作、833 个采样时间点；文件校验通过。')
    bullet('用户已认可机器人 09:40:37 版本的整体行走与关节衔接，脚与腿的前后遮挡仍需调整；其他版本和角色继续单独验收。')

    # 2 Function overview.
    page('01 平台功能与使用范围')
    p('平台提供三条制作路线。已有 Spine 换皮主要修改贴图；骨骼与贴图入口允许改动角色结构；单图入口从一张角色图构建新的基础人形动画。输入条件、输出格式与验收方式分别不同，接收方应按资源现状选择。[S1]')
    figure(ASSETS/'routes.png','三条制作路线及主要输出 步骤示意')
    table(['模块','当前可做的内容','主要边界'],[
        ('贴图换皮','部件或图集改色、上传替换、AI 材质生成、保留动作导出','不能直接添加原轮廓以外的配饰'),
        ('骨骼与贴图','调整骨骼、插槽、图片附件和双骨 IK；重打图集','修改结构后需逐动作检查'),
        ('单图行走','关节定位、拆件、补全、组装校准、idle 与 walk','适合四肢可辨的人形双足角色'),
        ('AI 与 Skill','根据风格和部件元数据生成提示词、提交任务和下载','缺少通用视觉语义分析和质量筛选'),
        ('历史与批量','保存恢复版本；按角色建立任务并汇总处理','缺少平台内的多角色批次总看板'),
        ('局域网部署','本机服务及局域网访问，供 Agent 远程调用','多人身份权限与运维能力仍需完善'),
    ],[1.08,3.46,2.34])
    p('本批多角色处理由独立脚本遍历目录并逐套调用 Skill 完成，属于已经验证的 Agent 编排方式；平台尚未提供通用的多角色一键上传、统一批次调度和管理页面。')

    # 3 Texture-only workflow.
    page('02 已有 Spine 的贴图换皮')
    p('支持导入 JSON 或 SKEL、atlas 和全部 PNG，或相应 ZIP；支持子目录、多页图集、裁切及 0／90／180／270 度图集旋转。预览入口覆盖 Spine 3.8 和 4.2。贴图模式沿用原骨骼、mesh 网格、蒙皮权重及动画曲线。[S1]')
    figure(ROOT/'evidence/editor.png','早期换皮工作台示例 选择身体部件调色并预览原版行走 [S1]',5.8)
    h('可操作的内容',2)
    bullet('选择单个或多个部件，下载原部件，调整色相、饱和度和明度。')
    bullet('上传同尺寸部件 PNG，或替换整张图集页；AI 可生成材质候选后再应用。')
    bullet('切换动作、皮肤、播放速度、骨骼显示和整体镜像；整体镜像同时改变脚掌朝向。')
    h('导出与保护范围',2)
    p('导出保持原 JSON／SKEL 和 atlas 字节，锁定贴图尺寸与 alpha 透明轮廓，并附文件校验报告。alpha 表示像素透明度；PMA 表示预乘透明度，需要遵循原图集声明。文件校验可以排除骨架被误改和透明轮廓漂移，不能判断五官、图案与材质是否合适。')
    p('已有贴图区域之外的结构变化，应转入骨骼与贴图编辑。原生 .spine 工程仍需先导出运行时资源，不能直接作为通用导入格式。')

    # 4 Rig editing and its distinction from retargeting.
    page('03 骨骼与贴图编辑及定位体验')
    p('从已有项目建立独立编辑方案后，可调整骨骼位置、旋转、缩放、长度、名称和父子关系，也可新增骨骼或删除符合引用约束的叶骨骼。新贴图可采用不同尺寸和透明度，并通过插槽与附件参数重新摆放。[S1][S7]')
    figure(ASSETS/'bone-inspection.png','骨骼定位示例 名称 关节点 骨长和受影响贴图可视化 [S7]',5.4)
    p('针对“移动时不知道在改哪个骨骼”的反馈，已增加选中骨骼名称、关节点、影响部件及子骨骼关系显示；可在画布选中、定位和拖动，支持未保存预览、应用与取消。右图保留当时的测试材质，仅用于说明 mesh 的影响范围。')
    h('骨架更换与动作重定向的区别',2)
    p('当前可重新绑定插槽、增删普通图片附件、编辑附件位置和旋转，调整双骨 IK 的目标与链条，并导入相同版本、使用当前图集的完整骨架 JSON。编辑后沿用现有局部动作曲线，比例和层级变化可能改变实际动作效果。')
    p('自动动作重定向是把来源角色的动作迁移到骨骼名称、长度或结构不同的目标角色，需要骨骼对应、比例补偿、曲线转换和接地约束。该通用能力尚未接入。此入口主要支持 3.8 和 4.2 JSON；SKEL 需要先转换为匹配版本的 JSON，输出也不生成原生 .spine 工程或新的 SKEL。')

    # 5 Single image workflow.
    page('04 根据一张图制作基础行走')
    p('支持上传 PNG、JPG 或 WebP，或通过亿创生成全身角色图。平台按轮廓比例给出建议关节点，可拖动或填写坐标；随后去背景、清理腿间背景、调整部件多边形，再建立人形双足骨架。[S1][S5]')
    figure(ROOT/'evidence/single-image/preserved-fill/comparison.png','原图保持路线的离线对比示例 用户认可的 09:40:37 版本另见第 07 节 [S5]',6.75)
    h('已具备的制作流程',2)
    bullet('拆分身体、背包、头和四肢；定位原图部件，检查邻近像素混入和连接情况。')
    bullet('使用角色参考图生成完整部件候选，默认只补遮挡区域，保留可见像素与原关节位置。')
    bullet('在组装校准页调整部件平移、旋转、缩放和前后层级，再生成动画版本。')
    bullet('生成 idle 与原地循环 walk，调整周期、步幅、抬脚高度和朝向，导出 Spine 4.2 JSON、atlas、图集、部件及关节配置。')
    h('输入条件决定结果上限',2)
    p('适合站立、四肢可辨且轮廓清晰的人形双足角色。建议关节点属于比例估计，不能等同于任意姿态识别。一张合成图没有被遮挡部位的完整像素；复杂姿态、背包遮挡和鞋口连接仍需补充素材或人工校准。')

    # 6 Single image optimization history.
    page('05 单图角色的制作与优化过程')
    p('机器人案例的反馈依次涉及关节撕裂、部件粘连、身体与背包混入、遮挡部位缺失，以及补绘后手臂和鞋子偏离原图。处理策略从连续蒙皮扩展到肢体隔离，再转向保留原图的遮挡补全和可逆组装校准。[S5][S6]')
    figure(ROOT/'evidence/single-image/seam-fix/comparison.png','撕裂修复过程 共享关节顶点和连续权重减少接缝 [S5]',6.75)
    table(['反馈问题','已采取的处理','仍需复验'],[
        ('关节撕裂','共享网格与连续权重','运动全周期是否露缝'),
        ('手脚粘连','身体与各肢体链分别蒙皮','交叉遮挡时能否独立分离'),
        ('身体背包混入','收紧拆分边界，定位与像素块检查','语义分层是否正确'),
        ('遮挡导致缺失','独立部件参考生成，只补遮挡像素','补全部位是否完整且自然'),
        ('原貌与鞋口变化','锁定可见像素与关节，调整鞋口和小腿绑定','比例、鞋跟与脚底接地'),
    ],[1.16,3.15,2.57])
    p('可见像素保留、循环误差为零和顶点连续是有价值的技术检查。后续用户已认可 09:40:37 版本的整体行走与关节衔接，剩余反馈集中在脚与腿的前后遮挡。该认可针对指定机器人版本；其他版本和角色仍应单独进行美术验收。')

    # 7 Assembly calibration.
    page('06 组装校准与行走检查')
    p('平台参考原项目的拆件摆放与可逆校准思路，独立实现“组装校准”页：先对照原图确认部件位置和层级，再生成行走，减少错误组装被直接带入动作。此页已具备部件选择、拖动、数值输入、原图叠加、单部件查看和复位。[S6]')
    figure(ROOT/'evidence/assembly-review/browser-assembly.jpg','实际组装校准页面 部件范围高亮 数值调整 层级和保存入口 [S6]',6.1)
    figure(ASSETS/'robot-walk.png','同一行走周期的四个代表姿态 用于检查摆臂 迈步和鞋口衔接 [S5]',6.3)
    p('校准配置随历史保存，导出增加 assembly-review.json。检查会提示断开的像素块和较大的补全面积；这些提示不能替代部件语义判断。相关记录显示未改部件顶点与骨骼曲线保持，并已验证页面加载、选择、高亮与排版；全部编辑动作的完整浏览器验收仍需扩展。')

    page('07 用户验收反馈与剩余调整')
    p('用户指定 2026 年 10 月 9 日 09:40:37 生成的行走预览，反馈“除了脚在腿前面，其他都可以，关节处也很棒”。历史记录对应“亿创完整部件 · 遮挡补全行走”，不是后续 10:01:00 的“保留原图 · 遮挡补全 · 鞋底稳定”版本。[用户反馈]')
    figure(ASSETS/'user-accepted-walk.png','用户提供的 09:40:37 版本截图 整体效果与关节衔接获认可 脚腿前后关系仍待调整',4.5)
    h('验收结论的更新',2)
    p('该机器人版本的整体行走、部件衔接与关节表现已获得用户认可。此前报告对单图行走使用笼统的“视觉未验收”表述，应收窄到具体版本和具体问题。原项目独立试验以及其他版本的未验收结论仍分别保留。')
    h('下一轮建议保留与调整的内容',2)
    bullet('以 09:40:37 历史版本作为对照基准，保留其已认可的部件、骨骼与蒙皮效果。')
    bullet('先核查近远侧脚、腿及鞋口的绘制顺序和遮挡范围；截图不足以确定唯一原因，不应直接重做全套骨架。')
    bullet('调整后逐帧查看两侧迈步、重叠和朝向切换，确认前后关系正确，同时保持关节衔接效果。')
    p('本条用户反馈作为具体案例的验收证据记录；本次报告更新没有重新生成图片或修改行走资源。')

    page('08 生图提示词与拆件质量')
    p('用户强调亿创提示词对后续拆件十分重要。提示词需要同时描述角色外观和拆件交付要求。本次截图中的生图虽然要求分离部件，头、躯干和腿仍保持连接，且额外出现手部；只写“sprite sheet”或“separated parts”不足以作为拆件验收依据。[用户反馈][S8]')
    figure(ASSETS/'user-matrix-prompt.png','用户提供的亿创生成截图 要求分离部件后仍出现连接身体和重复手部',5.4)
    table(['提示词要求','建议写法与检查目标'],[
        ('角色一致性','依据同一参考图，保持五官、比例、颜色、描边与光照，不重新设计角色。'),
        ('部件数量与布局','指定 16 个部件、固定 4×4 布局，每个格子只放一个部件；格子位置定义部件身份。'),
        ('左右和朝向','左右以角色自身为准，保持参考视角；两侧手脚分别生成，不用复制镜像替代。'),
        ('遮挡补全','背包、上臂、小腿及鞋子完整呈现；补出隐藏轮廓和关节连接余量，不裁断连接端。'),
        ('清晰分离','各部件留足空白，不连接、不重叠；背景统一纯色，无场景、接地阴影和文字水印。'),
        ('无额外内容','只生成清单中的部件，不出现完整站立角色、额外手脚、装饰或其他人物。'),
    ],[1.37,5.51],small=10,compact=True)
    h('可试用的提示词骨架',2)
    p('依据参考图制作同一角色的 2D 骨骼动画部件图。使用固定 4×4 布局，每格一个独立完整部件。第 1 行：头、背包、躯干、骨盆；第 2 行：角色左上臂、左前臂、左手、右上臂；第 3 行：右前臂、右手、左大腿、左小腿；第 4 行：左脚、右大腿、右小腿、右脚。保持原设计与参考视角，完整补出遮挡区域，并保留关节连接余量。部件之间留足纯白空隙，不连接、不重叠。不要完整人物、重复部件、文字、额外装饰或场景。')
    p('该模板尚未试验，本次未新增生图调用。生成后核对数量、身份与完整性；若仍连接或重复，转为逐部件生成并沿用统一参考和连接端要求。纯色背景需要去除，提示词不保证透明 alpha。')

    # Independent original-project experiment remains a separate conclusion.
    page('09 原项目的独立验证与结论')
    p('已独立部署 GenielabsOpenSource/spine-animation-ai，用同一张机器人原图验证拆分与行走。原 UI、拆分、定位和动画脚本保持原样；本轮生图传输接入现有亿创，生成 1 张图片，保留原提示词。该试验包含供应商适配与 Agent 辅助，不能视为原生 Gemini 的端到端质量结论。[S8]')
    figure(ASSETS/'upstream-parts.png','原生成图连通域只能拆出 5 块 Agent 多边形细分后形成 16 个部件 [S8]',6.2)
    figure(TRIAL/'site/media/upstream-walk-frames.png','原项目 walk 预设的四个采样姿态 可见关节切边和较大的脚部转动 [S8]',6.2)
    p('原 SIFT 定位在辅助拆件后可定位 16 块，原构建脚本输出 17 根骨骼、16 个部件；idle 与 walk 各 49 个采样点均无非有限顶点，循环误差为零。但实际仍可见手臂错位、背包缺边、鞋底缺损及缺少支撑脚落地约束，视觉未通过验收。')
    p('建议保留独立试验，选择性借鉴定位、校准与预设结构；优先改善完整部件生成和遮挡补全。直接整体替换平台尚不能解决本轮发现的问题。原项目及运行库随附的许可文件应保留。')

    # 9 AI/Skill/API.
    page('10 AI 调用与任务持久化')
    p('spine-ai-reskin Skill 可接收 Spine 目录、ZIP 或已有项目编号及目标风格。服务依据部件名称、尺寸、原色统计和风格模板组装提示词，先返回计划和生成数量，再按授权范围调用亿创。当前没有额外的视觉语义模型识别或 LLM 提示词优化步骤。[S2]')
    figure(ASSETS/'automation.png','AI 任务调用与恢复流程 编号复用用于避免盲目重复提交')
    table(['接口或工具','用途'],[
        ('GET /api/automation/openapi.json','读取当前调用规范'),
        ('POST /api/automation/reskin/plan','核对部件、实际提示词与数量，不提交生图'),
        ('POST /api/automation/reskin/tasks','建立独立结果项目并提交任务'),
        ('tasks/{id} 及 download resume cancel','查询结果、下载和按状态恢复或取消后续处理'),
        ('spine_agent.py 与 spine-ai-reskin','Agent 使用 plan、run、status、wait、download'),
    ],[3.37,3.51])
    h('生成路线和边界',2)
    p('默认 material 路线生成表面材质，再贴回原轮廓与明暗。reference 路线需要每个部件对应的公共 HTTPS 参考画布，当前不自动上传源素材。默认最多 32 次生成，每个部件一次；本批小厨神 43 个部件按“全部处理”范围显式配置上限。')
    p('状态、nonce 和已有生成编号持久保存。等待超时可继续查询同一任务；提交结果不明时先核实，不能换编号盲目重发。失败任务保留候选和原因，授权重做后才能建立新任务。凭据由服务器管理，不传给调用方，也不进入导出资源包。')

    # 10 Quantitative case.
    page('11 冰雪蓝白批量制作实测')
    p('输入为 Flambé 项目目录内 8 套 Spine 3.8.99 资源。先扫描全部部件、冻结输入与哈希，逐套生成计划，再复用独立任务编号提交和收集结果。目标为冰蓝、浅蓝和雪白，用户明确选择保留原肤色与五官。[S3]')
    figure(ASSETS/'batch-counts.png','本批请求与结果数量 成功候选数量不等于美术合格数量 [S3]',6.75)
    report=json.loads((BATCH/'final-report.json').read_text());batch=json.loads((BATCH/'batch.json').read_text())
    table(['资源','部件','成功','失败','未提交','AI 整套状态'],[
        (a['name'],e['parts'],a['job_counts'].get('completed',0),a['job_counts'].get('failed',0),a['job_counts'].get('prepared',0),'完成' if a['automation_state']=='completed' else '失败') for a,e in zip(report['assets'],batch['entries'])
    ],[1.28,.68,.68,.68,.86,2.70])
    p('60 个成功候选占实际提交 73 次的 82.2%，这是供应商返回图片的比例，不是视觉通过率。13 个部件失败，另 16 个小厨神部件在任务停止后没有提交。每个失败角色至少一个部件查到了亿创的安全策略拒绝；具体触发项未返回，不能据此确认是否误判。')
    p('3 套自动任务完整完成，5 套保持 failed。最终 8 套成品均采用成功 AI 候选配色与原图明暗拟色，失败或未提交部件用共同配色补齐；本地处理没有新增生图调用，也未将失败任务改成完成。')

    # 11 Gallery.
    page('12 八套最终换皮效果')
    figure(BATCH/'evidence/all-comparison.jpg','每组左侧为原角色 右侧为最终蓝白资源 人物身份与原动作保留 [S3]',5.72)

    # 12 Material failures and local fitting.
    page('13 生图纹样污染与本批清理方法')
    p('部分 AI 材质样本包含额外人脸和机械纹样。原有 material 应用方式会把样本空间图案缩放贴回部件，导致脸部覆盖多余五官、托盘出现人脸以及衣服图案失控。原色深浅被保留也不能防止这些内容进入浅色区域。[S3]')
    figure(ASSETS/'identity-comparison.png','原角色 清理前 AI 结果与最终拟色版 中间来自用户反馈截图 姿态时间不完全一致',5.8)
    p('本批改为只提取候选的蓝白配色，再根据原图明暗重映射。脸部和发型保留，四肢与领口保护肤色，不把候选里的空间纹样贴回。该方式更接近保持原造型的配色换皮，细腻霜纹或高保真新细节并未普遍实现。')
    figure(ASSETS/'uv-overlap.png','小厨神下半身与左上臂的 atlas 矩形范围交叠 按原 mesh UV 隔离可避免覆盖邻近部件 [S3]',5.8)
    p('UV 表示贴图采样坐标。部分 mesh 的矩形包围范围交叠，按矩形替换会把邻近部件一起处理。本批拟色脚本已按原 mesh UV 三角形限定写入，并保留少量边缘采样范围。这项处理及纹样清理属于本批交付脚本，仍需正式接入平台通用候选应用和自动交付流程。')

    # 13 Working project synchronization.
    page('14 工作台结果同步与历史管理')
    p('本批先完成独立交付包与对比页，工作台里的原自动结果尚未更新。用户随后打开旧项目，看到清理前的人脸纹样。问题来自两份结果的版本不同步；仅创建新的对比页面不足以保证原工作台展示最终成果。[S4]')
    figure(BATCH/'evidence/workbench-corrected.png','同步后的主角工作台 已恢复原五官 原项目和最终结果并排显示 [S4]',6.65)
    h('本次已经完成的修正',2)
    bullet('通过平台 API 同步 8 套最终图集，先保留旧版本，再保存可恢复的交付最终版。')
    bullet('使用 revision 版本号和原文件哈希检查，避免覆盖未记录的用户新编辑。')
    bullet('核对当前预览和工作台导出 PNG 与交付 PNG 字节一致，JSON、atlas、meta 保持原内容。')
    h('仍需进入产品流程的改进',2)
    p('AI 原始候选、已应用版本和交付最终版需要明确标识；预览、导出和交付应引用同一份最终快照。原自动任务的冻结 ZIP 保留用于审计，内容仍是当时的原始自动结果；当前工作台导出与本批最终交付包才是修正版，平台应避免让接收方混淆这两类下载。')

    # 14 Issues: statuses are carefully limited.
    page('15 当前不足及影响')
    p('主要问题集中在生成内容、拆件与遮挡、步态适配和最终版本管理。以下状态区分“本次案例已修正”和“通用平台能力仍待完善”，避免把一次人工辅助交付当作自动化能力已经完成。')
    table(['问题','表现及影响','当前状态'],[
        ('生图可靠性','安全审核拒绝；生成图混入多余人物或图案','本批已拟色补齐；自动质量筛选待接入'),
        ('角色一致性','脸、手臂、鞋和服装细节可能偏离原图','当前可保留原可见像素；隐藏部分仍需审核'),
        ('部件拆分','身体与背包混入；不同部件仍连接或缺块','边界编辑与组装校准已实现；自动语义拆件不足'),
        ('遮挡补全','背包、手臂、脚的被挡部分缺失','有独立部件参考生成；稳定性仍未验收'),
        ('蒙皮与步态','通用角色可能出现撕裂、穿插和脚底滑动','机器人 09:40 版获认可；脚腿遮挡仍待调整，其他角色需复验'),
        ('动作迁移','不同骨架之间难以直接沿用原动作','通用自动动作重定向尚未接入'),
        ('批量与版本','失败停止未提交部件；成品与预览曾不同步','本批脚本与同步已完成；统一批次看板和流程待建设'),
        ('部署与格式','局域网多人访问管理有限；原生工程和转换有限','需要身份权限、备份及匹配版本的转换工具链'),
    ],[1.15,2.71,3.02])
    h('验证结果的限制',2)
    p('17 个动作共 833 个采样时间点与原资源骨骼、绘制顶点一致；PNG 尺寸与 alpha 不变，原 JSON、atlas 和 meta 字节一致，原目录 40 个文件哈希不变。浏览器已切换全部 17 个动作，未记录播放器错误。这些结果证明资源兼容与动作保持，不代表每个像素、每个时刻或所有目标游戏均已通过美术验收。[S3]')
    p('当前也没有可据以承诺的通用高质量成功率、多人并发吞吐量或生图成本估计。平台应建立独立的视觉验收和运行监控数据，再确定这些指标。')

    # 15 Suggested improvements.
    page('16 建议的改善顺序与验收目标')
    p('建议优先把本批有效处理接入通用流程，随后改善单图素材质量和步态，最后扩展不同骨架间的动作迁移。下列内容是建议路线与验收要求，未给出开发工期或承诺完成日期。')
    table(['优先级','建议改进','建议验收方式'],[
        ('P0','把配色拟色、肤色保护和 mesh UV 写入隔离接入通用换皮','固定测试角色检查无邻近部件污染，五官、alpha 和动作按策略保留'),
        ('P0','增加候选质量检查和人工确认入口','禁止未检查候选直接成为最终版；能标记多余人脸、错部件和风格偏差'),
        ('P0','统一最终快照及预览导出；明确原始候选和审计冻结包','网页、下载和报告都能核对同一最终版本，历史可恢复'),
        ('P1','独立补全部件，锁定可见像素；改进拆分与前后层级','身体背包、左右手臂、腿和鞋分别审查，摆动后无明显缺块或粘连'),
        ('P1','增加接地约束、鞋口衔接和权重调整工具','全周期逐帧检查支撑脚、关节接缝、穿插和循环衔接'),
        ('P1','建设批次总看板、部件预算和失败分类恢复','单个失败不丢失其他结果；查询与重做分开；不重复提交已成功部件'),
        ('P2','先支持相近人形骨架的动作迁移，再扩大类型','骨骼映射可检查；比例补偿、脚底接地和新旧动作对照通过'),
        ('P2','完善身份权限、备份恢复与版本兼容评测','跨设备访问复验；项目隔离与恢复有效；逐版本说明输入输出范围'),
    ],[.62,3.06,3.20])
    h('生成失败应如何处理',2)
    p('区分内容审核、网络问题、未知提交和无图片输出。查询失败只恢复已有任务，内容审核应读取平台原因、调整任务范围或采用本地配色方案；需要再次生图时先确认授权，不绕过审核，也不因单个失败重复提交整套。预算以部件数、已提交数和待处理数分别显示。')

    # 16 Acceptance and practical delivery.
    page('17 交付使用与复验建议')
    h('已交付内容',2)
    bullet('工作台源码与启动入口，已有 Spine 换皮、骨骼编辑和单图角色制作流程。')
    bullet('OpenAPI、Python CLI、spine-ai-reskin Skill，及各模块的示例与验证记录。')
    bullet('Flambé 8 套蓝白资源 ZIP、逐套包、对比图、任务记录、校验与拟色说明。')
    bullet('原项目独立测试环境和机器人对照资源，保留原实现与辅助步骤的差别。')
    h('接收方建议复验顺序',2)
    table(['步骤','检查内容'],[
        ('1 导入与编辑','选择标准资源，确认 atlas 页齐全、部件和皮肤可识别；试改一块并保存历史。'),
        ('2 预览与动作','切换全部动作、皮肤和朝向，检查五官、接缝、穿插及脚底接地。'),
        ('3 导出与版本','下载当前项目，核对最终版本标记；在目标游戏中加载并确认贴图尺寸和动作兼容。'),
        ('4 AI 小规模试用','先查看 plan、提示词和预算，再按授权提交；中断后只查询原任务编号。'),
        ('5 单图视觉验收','逐部件对照原图，检查隐藏区域、背包层级、鞋口和完整行走周期。'),
        ('6 移交与运维','备份资源与历史，单独管理凭据；保留运行库和参考项目随附的许可文件。'),
    ],[1.18,5.70])
    h('本次部署入口',2)
    hyperlink('工作台  本机 8780','http://127.0.0.1:8780/')
    hyperlink('8 套蓝白资源对比与下载  本机 8895','http://127.0.0.1:8895/')
    hyperlink('原项目独立对照  本机 8894','http://127.0.0.1:8894/')
    p('服务监听地址为 0.0.0.0，其他电脑访问时使用服务器实际局域网 IP，不能把 0.0.0.0 当远端地址。本机地址仅对服务器所在电脑有效；服务进程和 IP 变化后需要重新确认入口。多用户环境还需要身份权限，固定操作请求头不能代替用户认证。')

    # 17 Traceable evidence appendix.
    page('附录 资料与验证记录')
    p('平台目录为 spine-reskin，原项目独立验证目录为 spine-upstream-trial，二者均位于 /Users/tu/Auto Creating。批量目录为 spine-reskin/delivery/flambe-ice-blue-white-20261009。以下编号对应正文的功能、截图和测量记录。')
    table(['编号','资料位置与用途'],[
        ('S1','spine-reskin/README.md、app/、web/\n平台功能、格式边界、单图与骨骼流程；较早截图用于说明当时的界面。'),
        ('S2','spine-reskin/agent/README.md\nskills/spine-ai-reskin/SKILL.md\nAgent 调用、生成路线、编号恢复与预算规则。'),
        ('S3','批量目录 batch.json、final-report.json\nevidence/runtime-validation.json\nevidence/browser-playback-checks.json\nevidence/package-validation.json\n8 套计数、拟色来源、17 动作和资源包校验。'),
        ('S4','批量目录 evidence/workbench-sync.json\nevidence/workbench-corrected.png\n8 套工作台同步、预览和导出一致性及最终截图。'),
        ('S5','spine-reskin/evidence/single-image/\nseam-fix/、preserved-fill/acceptance.json\n撕裂修正、原图保持、遮挡补全与运行时记录。'),
        ('S6','spine-reskin/evidence/assembly-review/\nacceptance.json、runtime-validation.json、browser-assembly.jpg\n组装校准、受影响部件及浏览器验证范围。'),
        ('S7','spine-reskin/evidence/rerig/\nbone-inspector-validation.json\n骨骼选择、高亮、影响部件及 3.8／4.2 示例。'),
        ('S8','spine-upstream-trial/README.md\nevidence/verification.json、raw-split.json、runtime-validation.json\n原项目独立验证、5 块拆分、16 部件辅助和视觉结论。'),
        ('用户反馈','用户提供的纹样污染、09:40:37 行走和亿创提示词截图\n本报告目录 user-feedback.json 记录指定版本及认可范围。'),
    ],[.74,6.14],small=9.5,compact=True)
    # Deterministic content and layout metadata for internal QA.
    doc.save(DOC_OUT)
    meta={'document':str(DOC_OUT),'figures':figure_number,'planned_pages':19,'source_cutoff':'2026-10-09',
          'new_ai_submissions':0,'existing_handover_unchanged':True,
          'scope':'平台功能与制作优化过程；真实批量结果；局部修正与平台待接入能力分开'}
    (OUT/'build-metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2))
    print(json.dumps(meta,ensure_ascii=False,indent=2))


if __name__=='__main__':build()
